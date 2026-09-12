import os
from openpyxl import Workbook, load_workbook

FILE_NAME = "student_results.xlsx"


# Create Excel file if it does not exist
def create_excel_file():
    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        headers = [
            "Roll No", "Name", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total", "Percentage", "Grade", "Status"
        ]

        sheet.append(headers)
        workbook.save(FILE_NAME)


# Calculate total, percentage, grade and status
def calculate_result(marks):
    total = sum(marks)
    percentage = total / 5

    # PASS only if every subject has at least 35 marks
    if all(mark >= 35 for mark in marks):
        status = "PASS"

        if percentage >= 90:
            grade = "A+"
        elif percentage >= 80:
            grade = "A"
        elif percentage >= 70:
            grade = "B"
        elif percentage >= 60:
            grade = "C"
        elif percentage >= 50:
            grade = "D"
        else:
            grade = "E"
    else:
        grade = "F"
        status = "FAIL"

    return total, percentage, grade, status


# Add a new student
def add_student():
    print("\n--------------------------------")
    print("       ADD STUDENT RESULT")
    print("--------------------------------")

    roll_no = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    student_class = input("Enter Course/Class: ")

    marks = []

    for i in range(1, 6):
        while True:
            try:
                mark = float(input(f"Enter marks for Subject {i}: "))

                if 0 <= mark <= 100:
                    marks.append(mark)
                    break
                else:
                    print("Marks must be between 0 and 100.")

            except ValueError:
                print("Please enter a valid number.")

    total, percentage, grade, status = calculate_result(marks)

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    # Check duplicate roll number
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == roll_no:
            print("A student with this Roll No already exists.")
            workbook.close()
            return

    sheet.append([
        roll_no,
        name,
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        grade,
        status
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    print("\nStudent result saved successfully!")
    print("--------------------------------")
    print(f"Total      : {total}")
    print(f"Percentage : {percentage:.2f}%")
    print(f"Grade      : {grade}")
    print(f"Status     : {status}")
    print("--------------------------------")


# Get result of a particular student
def get_result():
    print("\n--------------------------------")
    print("         STUDENT RESULT")
    print("--------------------------------")

    roll_no = input("Enter Roll No: ")

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == roll_no:
            found = True

            print("--------------------------------")
            print("Student Result")
            print("--------------------------------")
            print(f"Roll No     : {row[0]}")
            print(f"Name        : {row[1]}")
            print(f"Class       : {row[2]}")
            print(f"Total       : {row[8]}")
            print(f"Percentage  : {row[9]:.2f}%")
            print(f"Grade       : {row[10]}")
            print(f"Status      : {row[11]}")
            print("--------------------------------")

            break

    workbook.close()

    if not found:
        print("Student with this Roll No was not found.")


# Show all student data
def show_all_data():
    print("\n==============================================================")
    print("                 ALL STUDENT DATA")
    print("==============================================================")

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    if sheet.max_row <= 1:
        print("No student records available.")
        workbook.close()
        return

    print(
        f"{'Roll No':<10}"
        f"{'Name':<15}"
        f"{'Class':<12}"
        f"{'Total':<10}"
        f"{'Percentage':<13}"
        f"{'Grade':<8}"
        f"{'Status':<8}"
    )

    print("-" * 76)

    for row in sheet.iter_rows(min_row=2, values_only=True):
        print(
            f"{str(row[0]):<10}"
            f"{str(row[1]):<15}"
            f"{str(row[2]):<12}"
            f"{str(row[8]):<10}"
            f"{row[9]:<13.2f}"
            f"{str(row[10]):<8}"
            f"{str(row[11]):<8}"
        )

    print("==============================================================")

    workbook.close()


# Menu function
def menu():
    create_excel_file()

    while True:
        print("\n========================================")
        print("      STUDENT RESULT MANAGEMENT")
        print("========================================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")
        print("========================================")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("\nThank you for using Student Result Management System!")
            break

        else:
            print("\nInvalid choice! Please enter 1, 2, 3 or 4.")


# Program starts here
if __name__ == "__main__":
    menu() 
